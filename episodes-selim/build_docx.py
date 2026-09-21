#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a nicely formatted RTL Arabic Word document of the Selim episode script."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = RGBColor(0x1B, 0x7A, 0x43)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x66, 0x66, 0x66)


def rtl_paragraph(p, right=True):
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)


def styled_run(p, text, size=14, bold=False, color=DARK):
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), 'Arial')
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(size * 2))
    rPr.append(szCs)
    rtl = OxmlElement('w:rtl')
    rtl.set(qn('w:val'), '1')
    rPr.append(rtl)
    return r


def add_par(doc, text='', size=14, bold=False, color=DARK, space_after=8):
    p = doc.add_paragraph()
    rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        styled_run(p, text, size=size, bold=bold, color=color)
    return p


def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def cell_text(cell, text, size=14, bold=False, color=DARK):
    p = cell.paragraphs[0]
    rtl_paragraph(p)
    styled_run(p, text, size=size, bold=bold, color=color)


def main():
    doc = Document()
    # Page base font
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(14)

    with open('/home/user/selim_table3/script_doc.txt', encoding='utf-8') as f:
        lines = [ln.rstrip('\n') for ln in f]

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('# '):
            t = add_par(doc, line[2:], size=22, bold=True, color=GREEN, space_after=10)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # subtle bottom border under title
            pPr = t._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '12')
            bottom.set(qn('w:space'), '4')
            bottom.set(qn('w:color'), '1B7A43')
            pbdr.append(bottom)
            pPr.append(pbdr)
        elif line.startswith('## '):
            add_par(doc, line[3:], size=16, bold=True, color=GREEN, space_after=6)
        elif line.startswith('ملاحظة') or line.startswith('الحلقة'):
            add_par(doc, line, size=12, color=GRAY, space_after=12)
        else:
            add_par(doc, line, size=14, space_after=10)

    # --- Multiplication table ---
    add_par(doc, '')
    add_par(doc, 'جدول ضرب ٣ — للمراجعة السريعة', size=16, bold=True, color=GREEN, space_after=8)

    ar_nums = ['٠', '١', '٢', '٣', '٤', '٥', '٦', '٧', '٨', '٩', '١٠',
               '١١', '١٢', '١٣', '١٤', '١٥', '١٦', '١٧', '١٨', '١٩', '٢٠',
               '٢١', '٢٢', '٢٣', '٢٤', '٢٥', '٢٦', '٢٧', '٢٨', '٢٩', '٣٠']

    results = [3, 6, 9, 12, 15, 18, 21, 24, 27, 30]
    table = doc.add_table(rows=11, cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT

    cell_text(table.rows[0].cells[1], 'المسألة', size=14, bold=True)
    cell_text(table.rows[0].cells[0], 'الناتج', size=14, bold=True)
    shade_cell(table.rows[0].cells[0], 'DFF3E4')
    shade_cell(table.rows[0].cells[1], 'DFF3E4')

    for i in range(1, 11):
        cell_text(table.rows[i].cells[1], f'٣ × {ar_nums[i]}', size=14)
        cell_text(table.rows[i].cells[0], ar_nums[results[i - 1]], size=14, bold=True, color=GREEN)

    # --- Tips for parents ---
    add_par(doc, '')
    add_par(doc, 'نصايح سريعة للأهل', size=16, bold=True, color=GREEN, space_after=6)
    tips = [
        '• جلسات قصيرة: دقيقتين أو تلاتة كل يوم أحسن من جلسة طويلة مرة في الأسبوع.',
        '• لو سليم غلط، قول له: «قرّبت قوي، جرّب تاني» — المديح أسرع طريق للحفظ.',
        '• اربط الجدول بحاجات يومية: تمر، مصروف، صلاة، لعب… زي الحلقة بالظبط.',
        '• اسألوه بالمقلوب كمان: «١٢ = ٣ × كام؟» عشان الفهم يتثبّت.',
    ]
    for tip in tips:
        add_par(doc, tip, size=13, color=DARK, space_after=4)

    out = '/home/user/selim_table3/سكريبت_حلقة_سليم_جدول_٣.docx'
    doc.save(out)
    print('saved:', out)


if __name__ == '__main__':
    main()
